"""Read-only reconciliation of author XLSX annotations with the pinned CSV mirror.

Run with the bundled document Python runtime (openpyxl); emit JSON only.
"""
import argparse
import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


def sha(path):
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def audit(original, mirror, output):
    results, extracted = [], {}
    for number in range(1, 11):
        name = f"V{number}"
        path = original / f"{name}.xlsx"
        workbook = load_workbook(path, read_only=True, data_only=False)
        sheet = workbook.worksheets[0]
        columns = (0, 2, 4) if number in (1, 9, 10) else (0, 1, 2)
        rows, skipped = [], []
        for excel_row, values in enumerate(sheet.iter_rows(max_row=10000, max_col=6, values_only=True), 1):
            a, b, c = (values[i] if i < len(values) else None for i in columns)
            if not any(v is not None for v in (a, b, c)) or excel_row == 1:
                continue
            if (not isinstance(a, (int, float)) or not isinstance(b, (int, float))
                    or not isinstance(c, str) or a != int(a) or b != int(b)):
                skipped.append({"excelRow": excel_row, "values": [a, b, c]})
                continue
            rows.append({"excelRow": excel_row, "startFrame": int(a), "endFrame": int(b),
                         "rawLabel": c, "label": "_".join(c.strip().lower().split())})
        workbook.close()
        with (mirror / "clean_annotation_data" / f"{name}.csv").open(encoding="utf-8-sig", newline="") as stream:
            csv_rows = [{"startFrame": int(r["start_frame"]), "endFrame": int(r["end_frame"]),
                         "label": "_".join(r["class"].strip().lower().split())} for r in csv.DictReader(stream)]
        differences = [{"index": i, "original": a, "mirror": b} for i, (a, b) in enumerate(zip(rows, csv_rows))
                       if any(a[k] != b[k] for k in ("startFrame", "endFrame", "label"))]
        results.append({"source": name, "originalSha256": sha(path),
            "mirrorSha256": sha(mirror / "clean_annotation_data" / f"{name}.csv"), "sheet": sheet.title,
            "originalRows": len(rows), "mirrorRows": len(csv_rows), "rowWiseDifferences": differences,
            "unparsedSourceRows": skipped, "exactNormalizedRowMatch": len(rows) == len(csv_rows) and not differences and not skipped})
        extracted[name] = rows
    report = {"status": "SOURCE_TABLE_RECONCILIATION", "scriptSha256": sha(Path(__file__)), "sources": results,
              "annotations": extracted, "limitations": ["Matching tables do not prove skeleton row alignment or label correctness.",
                  "Formula/helper columns are not used as annotation values.", "Labels and original files are not changed."]}
    with output.open("x", encoding="utf-8") as stream:
        json.dump(report, stream, indent=2, allow_nan=False)
    print(json.dumps([{k: r[k] for k in ("source", "originalRows", "mirrorRows", "exactNormalizedRowMatch")}
                      | {"differences": len(r["rowWiseDifferences"]), "unparsed": len(r["unparsedSourceRows"])} for r in results], indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--original", type=Path, required=True)
    parser.add_argument("--mirror", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    audit(args.original, args.mirror, args.output)
